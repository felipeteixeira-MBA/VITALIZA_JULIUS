# HISTÓRICO: não é pipeline canônico; verREADME eCHANGELOG antes de executar.
from pathlib import Path
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy.stats import norm, chi2_contingency
from docx import Document
from docx.shared import Inches
import json, zipfile

OUT=Path('entregas_vitaliza');OUT.mkdir(exist_ok=True)
F={p.stem:pd.read_csv(p,low_memory=False) for p in Path('.').glob('*.csv')}
def dt(s):
 s=s.astype('string').str.strip();z=s.str.contains(r'Z$|[+-]\d\d:\d\d$',na=False)
 o=pd.Series(pd.NaT,index=s.index,dtype='datetime64[ns]')
 o.loc[z]=pd.to_datetime(s[z],format='mixed',errors='coerce',utc=True).dt.tz_convert('America/Sao_Paulo').dt.tz_localize(None)
 o.loc[~z]=pd.to_datetime(s[~z],format='mixed',errors='coerce')
 return o
def money(s):
 return pd.to_numeric(s.astype(str).str.replace('R$','',regex=False).str.replace(' ','',regex=False).str.replace(',','.',regex=False),errors='coerce')
E=F['eventos_assinatura'].copy();E['dt']=dt(E.data_hora);E=E.sort_values('dt')
A=F['assinaturas'].copy();A['s']=pd.to_datetime(A.inicio);A['t_snapshot']=pd.to_datetime(A.data_fim)
P=F['pagamentos'].copy();P['dt']=dt(P.data_hora);P['valor_num']=money(P.valor)
U=F['usuarios'].copy();US=F['uso_mensal'].copy();US['month']=pd.to_datetime(US.mes+'-01')
ACT=pd.read_parquet('atividades.parquet');ACT['dt']=dt(ACT.inicio)
expected={'usuarios':64948,'assinaturas':32643,'amizades':113480,'campanha_ago26':14120,'eventos_assinatura':275565,'pagamentos':165548,'uso_mensal':591662,'onboarding':63929,'pesquisa_saida':643,'feed_posts':119309,'tickets_suporte':4422,'atividades':1254475}
ids={'usuarios':['user_id'],'assinaturas':['assinatura_id'],'amizades':['seguidor_id','seguido_id'],'campanha_ago26':['user_id'],'pagamentos':['pagamento_id'],'uso_mensal':['user_id','mes'],'onboarding':['onboarding_id'],'pesquisa_saida':['resposta_id'],'feed_posts':['post_id'],'tickets_suporte':['ticket_id'],'atividades':['atividade_id']}
quality=[];nulls=[];orphans=[]
for name,d in dict(F,atividades=ACT.drop(columns='dt')).items():
 quality.append(dict(base=name,linhas=len(d),esperado=expected[name],confere=len(d)==expected[name],duplicatas_exatas=int(d.duplicated().sum()),duplicatas_chave=int(d.duplicated(ids[name]).sum()) if name in ids else np.nan))
 for c in d:
  nulls.append(dict(base=name,coluna=c,nulos=int(d[c].isna().sum()),pct_nulos=100*d[c].isna().mean()))
 for c in ['user_id','seguidor_id','seguido_id','treino_com']:
  if c in d:
   v=d[c].dropna();orphans.append(dict(base=name,coluna=c,referencias=len(v),orfaos=int((~v.isin(U.user_id)).sum())))
quality=pd.DataFrame(quality);nulls=pd.DataFrame(nulls);orphans=pd.DataFrame(orphans)
checks=[]
def check(name,n,action):checks.append(dict(verificacao=name,ocorrencias=int(n),tratamento=action))
check('Datas de eventos inválidas',E.dt.isna().sum(),'Normalização UTC para Brasília; horários sem fuso tratados como Brasília.')
check('Valores monetários com R$',P.valor.astype(str).str.contains('R$',regex=False).sum(),'Conversão para número; preservar coluna original.')
check('Pagamentos negativos',P.valor_num.lt(0).sum(),'Preservar e interpretar por status; não contar como cobrança positiva.')
check('Valor monetário inválido após conversão',P.valor_num.isna().sum(),'Investigar antes de agregar receita.')
check('Contas com sistema fora da grafia canônica',(~U.sistema.isin(['android','ios'])).sum(),'Normalizar minúsculas; não usar versão atual em cortes históricos sem histórico.')
check('Cortesias com status ativo',((A.plano_atual=='cortesia')&(A.status_atual=='ativa')).sum(),'Excluir da base pagante.')
check('Usuários com múltiplas assinaturas neste arquivo',A.user_id.duplicated().sum(),'Nesta entrega há uma assinatura por usuário; não assumir isso para dados futuros.')
check('Treinos com data inválida',ACT.dt.isna().sum(),'Excluir apenas de agregações temporais se inválido; manter para auditoria.')
check('Treinos com duração não positiva',pd.to_numeric(ACT.duracao_min,errors='coerce').le(0).sum(),'Sinalizar; não corrigir automaticamente.')
check('Uso com dias ativos acima dos dias do mês',US.dias_ativos.gt(US.month.dt.days_in_month).sum(),'Sinalizar inconsistência.')
check('Uso com soma dos tipos diferente do total',US.atividades.ne(US.atividades_ar_livre+US.atividades_video).sum(),'Não reconciliar por imputação automática.')
bday=pd.to_datetime(U.data_nascimento,format='mixed',errors='coerce')
check('Nascimento não vazio mas inválido',(U.data_nascimento.notna()&bday.isna()).sum(),'Tratar como desconhecido.')
check('Nascimento futuro no corte',bday.gt('2026-09-30').sum(),'Tratar como desconhecido; não usar idade sem validar.')
checks=pd.DataFrame(checks)
# Paid intervals: events are authoritative for starts/ends; snapshot supplies legacy start only.
B=A[A.plano_atual.isin(['mensal','anual'])].copy()
starts=E[E.evento.eq('assinatura_iniciada')].set_index('user_id').dt.dt.normalize()
ends=E[E.evento.eq('assinatura_expirada')].set_index('user_id').dt.dt.normalize()
B['inicio_fonte']=np.where(B.user_id.isin(starts.index),'evento','cadastro_legado')
B['s']=B.user_id.map(starts).fillna(B.s);B['t']=B.user_id.map(ends)
plan_events=E[E.plano.isin(['mensal','anual'])].copy()
flags=E[E.evento.isin(['renovacao_automatica_desligada','renovacao_automatica_religada'])].copy()
flags['renovacao_desligada']=flags.evento.eq('renovacao_automatica_desligada')
def active(t):return B[(B.s<t)&(B.t.isna()|B.t.ge(t))].copy()
def enrich(b,t):
 last=plan_events[plan_events.dt<t].drop_duplicates('user_id',keep='last').set_index('user_id')
 b['plano']=b.user_id.map(last.plano);b['canal']=b.user_id.map(last.canal_cobranca)
 # Legacy plan may be unknown before first event. Explicit unknown, not current snapshot.
 b['plano']=b.plano.fillna('desconhecido');b['canal']=b.canal.fillna('desconhecido')
 fl=flags[flags.dt<t].drop_duplicates('user_id',keep='last').set_index('user_id').renovacao_desligada
 b['renovacao_desligada']=b.user_id.map(fl).astype('boolean')
 b['tempo_dias']=(t-b.s).dt.days
 b['faixa_tempo']=pd.cut(b.tempo_dias,[0,30,90,180,365,np.inf],labels=['01_ate_30d','02_31a90d','03_91a180d','04_181a365d','05_mais365d'],include_lowest=True).astype(str)
 b['coorte']=b.s.dt.to_period('Q').astype(str)
 b=b.merge(U[['user_id','canal_aquisicao']],on='user_id',how='left',validate='one_to_one')
 return b
months=pd.date_range('2025-01-01','2026-09-01',freq='MS');series=[];panel=[]
for t in months:
 n=t+pd.offsets.MonthBegin(1);b=enrich(active(t),t);b['mes']=t.strftime('%Y-%m');b['perda']=b.t.ge(t)&b.t.lt(n)
 new=B[B.s.ge(t)&B.s.lt(n)];total_end=B.t.ge(t)&B.t.lt(n)
 series.append(dict(mes=t.strftime('%Y-%m'),base_inicial=len(b),novos=len(new),reativacoes=0,perdas_base_inicial=int(b.perda.sum()),saidas_entrantes_mes=int((new.t.ge(t)&new.t.lt(n)).sum()),fins_total=int(total_end.sum()),base_final=len(active(n)),churn=b.perda.mean()))
 panel.append(b)
S=pd.DataFrame(series);H=pd.concat(panel,ignore_index=True)
S['saldo_confere']=S.base_inicial+S.novos-S.fins_total==S.base_final
cut=enrich(active(pd.Timestamp('2026-10-01')),pd.Timestamp('2026-10-01'))
recon=pd.DataFrame([
 ['Pagantes em 30/09/2026',len(cut),14753],['Base inicial jan/2025',int(S.iloc[0].base_inicial),8540],
 ['Base inicial abr/2026',int(S.loc[S.mes.eq('2026-04'),'base_inicial'].iloc[0]),14237],
 ['Perdas abr/2026',int(S.loc[S.mes.eq('2026-04'),'perdas_base_inicial'].iloc[0]),1154],
 ['Base inicial set/2026',int(S.iloc[-1].base_inicial),14319],
 ['Fins jun-set/2026',int(S.loc[S.mes.ge('2026-06'),'fins_total'].sum()),3389],
 ['Perdas base inicial jun-set/2026',int(S.loc[S.mes.ge('2026-06'),'perdas_base_inicial'].sum()),3362]],columns=['metrica','calculado','enunciado'])
recon['confere']=recon.calculado.eq(recon.enunciado)
def grouped(cols):
 o=H.groupby(['mes']+cols,dropna=False,observed=True).agg(base_inicial=('user_id','size'),perdas=('perda','sum')).reset_index();o['churn']=o.perdas/o.base_inicial;return o
recortes={k:grouped([k]) for k in ['plano','canal','faixa_tempo','coorte','canal_aquisicao']}
# Comparable Jan-Sep annual summaries, retaining person-month denominator.
C=H[H.mes.str[5:7].le('09')].copy();C['ano']=C.mes.str[:4]
compare=C.groupby(['ano','plano','canal'],observed=True).agg(exposicoes=('user_id','size'),perdas=('perda','sum')).reset_index();compare['churn_ponderado']=compare.perdas/compare.exposicoes
# Prior completed month use/payment signals (never current outcome month).
prev=US.copy();prev['mes']=(prev.month+pd.offsets.MonthBegin(1)).dt.strftime('%Y-%m')
sig=H.merge(prev[['user_id','mes','atividades','dias_ativos','sessoes_app']],on=['user_id','mes'],how='left',validate='one_to_one')
sig['faixa_uso']=pd.cut(sig.atividades,[-np.inf,0,3,7,np.inf],labels=['zero','1a3','4a7','8mais']).astype('string').fillna('sem_registro')
P['mes']=(P.dt.dt.to_period('M').dt.to_timestamp()+pd.offsets.MonthBegin(1)).dt.strftime('%Y-%m')
decl=P[P.status.eq('recusado')].groupby(['user_id','mes']).size().rename('recusas_mes_anterior').reset_index()
sig=sig.merge(decl,on=['user_id','mes'],how='left');sig['recusa_anterior']=sig.recusas_mes_anterior.fillna(0).gt(0)
sig['renovacao_grupo']=sig.renovacao_desligada.astype('string').fillna('desconhecido')
support=F['tickets_suporte'].copy();support['dt']=dt(support.aberto_em)
support['mes']=(support.dt.dt.to_period('M').dt.to_timestamp()+pd.offsets.MonthBegin(1)).dt.strftime('%Y-%m')
tk=support.groupby(['user_id','mes']).size().rename('tickets_mes_anterior').reset_index();sig=sig.merge(tk,on=['user_id','mes'],how='left');sig['ticket_anterior']=sig.tickets_mes_anterior.fillna(0).gt(0)
# Previous-month activity change.
prev2=US[['user_id','month','atividades']].copy();prev2['mes']=(prev2.month+pd.offsets.MonthBegin(2)).dt.strftime('%Y-%m');prev2=prev2.rename(columns={'atividades':'atividades_menos2'})
sig=sig.merge(prev2[['user_id','mes','atividades_menos2']],on=['user_id','mes'],how='left')
sig['queda_50pct']=np.where(sig.atividades_menos2.gt(0)&sig.atividades.notna(),np.where(sig.atividades.le(sig.atividades_menos2*.5),'queda_50pct_ou_mais','sem_queda_50pct'),'nao_comparavel')
signal_tables={}
for col in ['faixa_uso','recusa_anterior','ticket_anterior','renovacao_grupo','queda_50pct']:
 g=sig.groupby(col,dropna=False,observed=True).agg(exposicoes=('user_id','size'),perdas=('perda','sum')).reset_index();g['churn']=g.perdas/g.exposicoes;signal_tables[col]=g
# Campaign: ITT 30 days using all randomized users, events strictly after dispatch.
camp=F['campanha_ago26'].copy().rename(columns={'canal':'canal_contato'});t0=pd.Timestamp('2026-08-10 10:00:00');t1=t0+pd.Timedelta(days=30)
pre=enrich(B.copy(),t0)[['user_id','plano','canal','tempo_dias','renovacao_desligada']]
camp=camp.merge(pre,on='user_id',how='left',validate='one_to_one')
terminal=E[E.evento.eq('assinatura_expirada') & E.dt.gt(t0)&E.dt.le(t1)]
camp['churn_30d']=camp.user_id.isin(terminal.user_id)
exact_s=B.user_id.map(E[E.evento.eq('assinatura_iniciada')].set_index('user_id').dt).fillna(B.s)
exact_t=B.user_id.map(E[E.evento.eq('assinatura_expirada')].set_index('user_id').dt)
eligible=B.loc[exact_s.le(t0)&(exact_t.isna()|exact_t.gt(t0)),'user_id']
camp['elegivel_reconstruido']=camp.user_id.isin(eligible)
july=US[US.mes.eq('2026-07')][['user_id','atividades','sessoes_app']]
camp=camp.merge(july,on='user_id',how='left',validate='one_to_one')
cg=camp.groupby('grupo').agg(n=('user_id','size'),perdas=('churn_30d','sum'),abriu_email=('abriu_email','sum'),clicou=('clicou','sum'),elegiveis=('elegivel_reconstruido','sum')).reset_index();cg['churn']=cg.perdas/cg.n
ctrl=cg[cg.grupo.eq('controle')].iloc[0];effect=[]
for r in cg.itertuples():
 if r.grupo=='controle':continue
 delta=ctrl.churn-r.churn;se=np.sqrt(ctrl.churn*(1-ctrl.churn)/ctrl.n+r.churn*(1-r.churn)/r.n);p=2*norm.sf(abs(delta/se)) if se else 1
 effect.append(dict(grupo=r.grupo,reducao_churn_pp=delta*100,ic95_inferior_pp=(delta-1.96*se)*100,ic95_superior_pp=(delta+1.96*se)*100,p_valor=p,p_bonferroni=min(1,p*3),retidos_incrementais_estimados=delta*r.n,custo_email=r.n*.01))
effect=pd.DataFrame(effect)
secondary=[]
for label,endpoint,restrict in [('ate_corte',pd.Timestamp('2026-10-01'),False),('30d_elegiveis',t1,True)]:
 tmp=camp[camp.elegivel_reconstruido].copy() if restrict else camp.copy()
 term=E[E.evento.eq('assinatura_expirada')&E.dt.gt(t0)&E.dt.lt(endpoint)]
 tmp['outcome']=tmp.user_id.isin(term.user_id)
 g=tmp.groupby('grupo').outcome.agg(['size','sum','mean']);c=g.loc['controle']
 for name,r in g.iterrows():
  if name=='controle':continue
  delta=c['mean']-r['mean'];se=np.sqrt(c['mean']*(1-c['mean'])/c['size']+r['mean']*(1-r['mean'])/r['size'])
  secondary.append(dict(analise=label,grupo=name,n=int(r['size']),perdas=int(r['sum']),churn=r['mean'],churn_controle=c['mean'],reducao_pp=delta*100,ic95_inf_pp=(delta-1.96*se)*100,ic95_sup_pp=(delta+1.96*se)*100,p_valor=2*norm.sf(abs(delta/se))))
secondary=pd.DataFrame(secondary)
strata=camp.groupby(['plano','grupo']).agg(n=('user_id','size'),perdas=('churn_30d','sum')).reset_index();strata['churn']=strata.perdas/strata.n
balance=camp.groupby('grupo').agg(n=('user_id','size'),tempo_medio_dias=('tempo_dias','mean'),atividades_julho_media=('atividades','mean'),sessoes_julho_media=('sessoes_app','mean'),pct_anual=('plano',lambda x:x.eq('anual').mean()),pct_lojas=('canal',lambda x:x.isin(['app_store','google_play']).mean())).reset_index()
balance_tests=[]
for col in ['plano','canal']:
 tab=pd.crosstab(camp.grupo,camp[col]);stat,p,_,_=chi2_contingency(tab);balance_tests.append(dict(variavel=col,teste='qui-quadrado',p_valor=p))
balance_tests=pd.DataFrame(balance_tests)
# A priori business-cost scenario; not observed offer redemption or realized ROI.
fees={'app_store':.15,'google_play':.15,'web_cartao':.035,'web_pix':.0099}
camp['margem_mes_proxy']=np.where(camp.plano.eq('anual'),299.9/12,39.9)*(1-camp.canal.map(fees))-4.2
econ=[]
for r in effect.itertuples():
 sub=camp[camp.grupo.eq(r.grupo)];offer=sub.oferta.map({'30pct_3_meses':35.91,'1_mes_gratis':25}).sum() if r.grupo=='cupom' else 0
 margin=sub.margem_mes_proxy.mean();gross=r.retidos_incrementais_estimados*margin
 econ.append(dict(grupo=r.grupo,retidos_incrementais=r.retidos_incrementais_estimados,margem_mensal_proxy=margin,contribuicao_incremental_1mes_proxy=gross,custo_email=r.custo_email,custo_ofertas_cenario_todos_recebem=offer,saldo_proxy_1mes=gross-r.custo_email-offer))
econ=pd.DataFrame(econ)
# Plots.
plt.style.use('seaborn-v0_8-whitegrid')
fig,ax=plt.subplots(figsize=(11,4));ax.plot(pd.to_datetime(S.mes),100*S.churn,marker='o',color='#4c3c78')
for date,label in [('2026-03-01','Reajuste'),('2026-06-01','Pesquisa'),('2026-08-10','Campanha')]:ax.axvline(pd.Timestamp(date),ls='--',alpha=.6,label=label)
ax.set(ylabel='Churn mensal (%)',title='Vitaliza — churn da base pagante inicial');ax.legend();fig.autofmt_xdate();fig.tight_layout();fig.savefig(OUT/'01_churn_mensal.png',dpi=160);plt.close(fig)
fig,ax=plt.subplots(figsize=(10,4))
for name,g in recortes['plano'].groupby('plano'):
 if name!='desconhecido':ax.plot(pd.to_datetime(g.mes),g.churn*100,label=name,marker='.')
ax.set(ylabel='Churn (%)',title='Churn por plano no início do mês');ax.legend();fig.autofmt_xdate();fig.tight_layout();fig.savefig(OUT/'02_churn_plano.png',dpi=160);plt.close(fig)
fig,ax=plt.subplots(figsize=(8,4));ax.bar(cg.grupo,cg.churn*100,color=['#7a6a99' if x=='controle' else '#42a599' for x in cg.grupo]);ax.set(ylabel='Churn em 30 dias (%)',title='Campanha 10/08 — intenção de tratar');fig.tight_layout();fig.savefig(OUT/'03_campanha.png',dpi=160);plt.close(fig)
# Export analytic products, with no emails or raw free text.
histcols=['user_id','assinatura_id','mes','s','t','plano','canal','tempo_dias','faixa_tempo','coorte','renovacao_desligada','perda']
H[histcols].to_parquet(OUT/'historico_mensal_assinantes.parquet',index=False)
B[['user_id','assinatura_id','s','t','inicio_fonte']].to_csv(OUT/'intervalos_assinaturas.csv',index=False)
cut[['user_id','assinatura_id','plano','canal','tempo_dias','renovacao_desligada']].to_csv(OUT/'pagantes_corte_2026_09_30.csv',index=False)
sig.drop(columns=['email'],errors='ignore').to_parquet(OUT/'sinais_anteriores_churn.parquet',index=False)
camp.to_csv(OUT/'campanha_resultados_individuais.csv',index=False)
tables={'qualidade':quality,'nulos':nulls,'integridade_referencial':orphans,'regras_qualidade':checks,'reconciliacao':recon,'serie_mensal':S,'comparacao_jan_set':compare,'campanha_grupos':cg,'campanha_efeitos':effect,'campanha_sensibilidade':secondary,'campanha_por_plano':strata,'campanha_balanceamento':balance,'testes_balanceamento':balance_tests,'economia_cenarios':econ}
tables.update({'churn_'+k:v for k,v in recortes.items()});tables.update({'sinal_'+k:v for k,v in signal_tables.items()})
with pd.ExcelWriter(OUT/'analises_vitaliza.xlsx',engine='openpyxl') as w:
 for k,v in tables.items():
  v.to_excel(w,sheet_name=k[:31],index=False)
  sh=w.sheets[k[:31]];sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
  for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=min(55,max(15,max(len(str(c.value or '')) for c in list(col)[:200])+2))
summary={'serie':S.to_dict('records'),'reconciliacao':recon.to_dict('records'),'campanha':cg.to_dict('records'),'efeitos':effect.to_dict('records'),'sinais':{k:v.astype(object).where(v.notna(),None).to_dict('records') for k,v in signal_tables.items()},'economia':econ.to_dict('records'),'qualidade':checks.to_dict('records'),'corte_planos':cut.plano.value_counts().to_dict()}
(OUT/'resultados_resumo.json').write_text(json.dumps(summary,ensure_ascii=False,default=str,indent=2))
print('RECONCILIACAO\n',recon.to_string(index=False));print('SERIE\n',S.to_string(index=False));print('CAMPANHA\n',cg.to_string(index=False));print('EFEITOS\n',effect.to_string(index=False));print('SINAIS');
for k,v in signal_tables.items():print(k,'\n',v.to_string(index=False))
print('QUALIDADE\n',checks.to_string(index=False));print('BALANCEAMENTO\n',balance.to_string(index=False));print('ECONOMIA\n',econ.to_string(index=False))
print('SENSIBILIDADE\n',secondary.to_string(index=False))
# Additional diagnostic tables and event-by-event state ledger.
exit_events=E[E.evento.eq('assinatura_expirada')].copy();exit_events['mes']=exit_events.dt.dt.strftime('%Y-%m')
exit_events['motivo_observado']=np.where(exit_events.detalhe.eq('falha_pagamento'),'falha_pagamento','outro_ou_nao_informado')
exit_reasons=exit_events.groupby(['mes','motivo_observado']).size().rename('fins_assinatura').reset_index()
month_compare=[]
for key in ['plano','canal','faixa_tempo','coorte','canal_aquisicao']:
 d=recortes[key];d=d[d.mes.str[5:7].le('09')].copy();d['ano']=d.mes.str[:4]
 g=d.groupby(['ano',key],observed=True)[['base_inicial','perdas']].sum().reset_index();g['churn_ponderado']=g.perdas/g.base_inicial;g=g.rename(columns={key:'grupo'});g.insert(0,'recorte',key);month_compare.append(g)
annual_groups=pd.concat(month_compare,ignore_index=True)
annual_summary=S.assign(ano=S.mes.str[:4],periodo=np.where(S.mes.str[5:7].le('09'),'jan_set','out_dez')).groupby(['ano','periodo']).agg(meses=('mes','size'),media_simples_churn=('churn','mean'),exposicoes=('base_inicial','sum'),perdas=('perdas_base_inicial','sum')).reset_index();annual_summary['churn_ponderado']=annual_summary.perdas/annual_summary.exposicoes
with pd.ExcelWriter(OUT/'analises_vitaliza.xlsx',engine='openpyxl',mode='a') as w:
 exit_reasons.to_excel(w,sheet_name='fins_por_motivo',index=False);annual_groups.to_excel(w,sheet_name='recortes_jan_set',index=False);annual_summary.to_excel(w,sheet_name='resumo_periodos',index=False)
state={};renew={};ledger=[]
for r in B[B.s.lt('2024-07-01')].itertuples():state[r.user_id]='pagante';ledger.append((r.user_id,pd.Timestamp('2024-07-01'),'saldo_inicial_legado','pagante',None))
for r in E.itertuples():
 old=state.get(r.user_id,'sem_assinatura');ev=r.evento
 if ev=='trial_iniciado':old='trial'
 elif ev=='trial_expirado' and old=='trial':old='sem_assinatura'
 elif ev in ['assinatura_iniciada','renovacao']:old='pagante'
 elif ev=='periodo_carencia':old='carencia'
 elif ev in ['assinatura_expirada','reembolso']:old='encerrada'
 if ev=='renovacao_automatica_desligada':renew[r.user_id]=False
 elif ev=='renovacao_automatica_religada':renew[r.user_id]=True
 state[r.user_id]=old;ledger.append((r.user_id,r.dt,ev,old,renew.get(r.user_id)))
L=pd.DataFrame(ledger,columns=['user_id','data_hora_brasilia','evento','situacao_apos_evento','renovacao_automatica_conhecida']);L.to_parquet(OUT/'historico_situacao_eventos.parquet',index=False)
# Written report in Portuguese.
def pct(x):return f'{x*100:.2f}%'.replace('.',',')
def num(x):return f'{int(x):,}'.replace(',','.')
lines=[]
def h(t):lines.extend(['',t,''])
def text(t):lines.append(t)
def table(d):
 lines.append(' | '.join(map(str,d.columns)))
 for row in d.itertuples(index=False,name=None):lines.append(' | '.join('' if pd.isna(v) else str(v) for v in row))
h('# Vitaliza — diagnóstico inicial e sete entregas')
text('Data de produção: 16/10/2026. Dados observados até 30/09/2026. Estudo baseado exclusivamente nos anexos; não utiliza dados externos. Nenhum modelo preditivo foi treinado nesta etapa.')
h('## 1. Resumo executivo')
text('A base paga foi reconciliada em 14.753 assinantes: 8.478 anuais e 6.275 mensais. Há ainda 60 cortesias ativas, excluídas do denominador de churn. Todos os sete pontos de conferência bateram e o saldo de entradas e saídas fecha nos 21 meses.')
text('A média simples das taxas mensais sobe de 5,62% em 2025 para 6,62% em jan-set/2026. Em períodos comparáveis, jan-set/2025 registra 5,61%. A subida já existe em janeiro de 2026, antes do reajuste de março; preço sozinho não explica todo o movimento.')
text('Nos recortes jan-set, a taxa mensal ponderada do plano mensal sobe de 9,14% para 10,86%; a do anual, de 2,49% para 3,32%. A comparação anual deve considerar a proximidade da renovação. Plano/canal históricos não identificáveis antes do primeiro evento válido ficam como desconhecidos, sem usar o retrato atual retroativamente.')
text('Uso baixo ou em queda, cobranças recusadas, suporte recente e renovação desligada estão associados a perdas posteriores. Essas relações são diagnósticas, não efeitos causais. As taxas agregam exposições usuário-mês, não pessoas distintas.')
text('No experimento, nenhuma mensagem apresenta diferença conclusiva em 30 dias. Até 30/09, cupom reduz o churn observado de 9,75% para 7,83%: diferença de 1,92 ponto percentual (IC95% de 0,60 a 3,24). É uma análise secundária exploratória, sem pré-registro; não demonstra retorno econômico positivo nem permite transportar o efeito para WhatsApp/CS.')
h('## 2. Relatório de qualidade')
text('Todos os 12 arquivos de dados, incluindo o Parquet, foram lidos. Contagens conferem com o dicionário. Detalhes de nulos, chaves e referências constam na planilha. Duplicatas exatas não substituem a análise de duplicatas de negócio.')
table(quality)
text('Principais verificações e tratamentos:');table(checks)
text('Os 3.649 treinos com duração não positiva são sinalizados, não apagados. As análises de uso desta etapa utilizam a tabela mensal do BI; não foi recalculada a duração mensal a partir desses treinos. A reconciliação detalhada BI versus treinos é pendência.')
text('Datas de nascimento futuras não devem produzir idades negativas no modelo. Nulos de GPS/frequência cardíaca podem ser estruturais. Campos atuais do cadastro e situação atual dos tickets não foram usados para reconstruir atributos históricos.')
text('A pasta de mapas de calor não foi recebida. Sua ausência não impede estas entregas, mas impede a avaliação das variáveis de imagem.')
h('## 3. Definições formais e regras de reconstrução')
text('Pagante: assinatura de plano mensal/anual iniciada e ainda não expirada; cortesias e trials excluídos. Carência é preservada no histórico de estados, sem encerrar automaticamente o ciclo pago. Chargeback isolado não é encerramento: exige evidência terminal.')
text('Início: evento assinatura_iniciada; para ciclos anteriores ao começo dos logs, início do cadastro de assinaturas. Fim efetivo: evento assinatura_expirada. Reembolso é marcado no histórico de situação, mas não gera segunda perda quando existe expiração do mesmo ciclo; os 72 reembolsos têm expiração associada.')
text('Cancelamento web e desligamento da renovação: intenção/decisão de não renovar, sem perda imediata do período vigente. Trial_expirado não é churn pago. Plano e canal: último evento com plano mensal/anual conhecido antes da referência; desconhecido quando não há evidência. A ausência de evento explícito de renovação automática fica desconhecida, não é assumida como ligada.')
text('Calendário mensal: início normalizado para o dia; base inicial inclui início anterior ao primeiro dia e fim no primeiro dia ou depois. Expirações no primeiro dia pertencem às perdas daquele mês. Base final é a base inicial do mês seguinte. Eventos são convertidos de UTC para Brasília antes de obter o dia. As convenções diárias reproduzem as métricas oficiais.')
text('Churn = perdas de ciclos presentes na base inicial / base inicial. Saídas de quem entrou no próprio mês ficam fora do numerador e entram na conciliação do saldo. Saldo = base inicial + entradas − todos os fins do mês. A base entregue contém uma assinatura por usuário e um evento de início por novo assinante: não há reativações identificadas. Zero reativações significa ausência identificada nos anexos, não regra geral do negócio.')
text('O histórico evento a evento registra trial, pagante, carência e encerramento. O painel mensal contém os pagantes no início de cada mês, não todos os usuários gratuitos. A definição de carência como acesso pago é operacional e precisa ser confirmada para uso em produção.')
h('## 4. Base histórica e reconciliação')
table(recon);text('Intervalos por assinatura, estados após cada evento, painel usuário-mês e pagantes no corte foram exportados. Não foram exportados e-mails, textos de saúde ou mensagens brutas nesses produtos.')
h('## 5. Série mensal de churn')
v=S[['mes','base_inicial','novos','perdas_base_inicial','saidas_entrantes_mes','base_final','churn']].copy();v['churn']=v.churn.map(pct);table(v)
text('Média simples de taxas mensais e taxa ponderada não são a mesma métrica. A planilha distingue as duas. As 3.389 expirações de jun-set incluem 27 saídas de entrantes do próprio mês; por isso a série de perdas da base inicial soma 3.362.')
h('## 6. Churn por plano, canal, tempo de casa e coorte')
view=annual_groups[annual_groups.recorte.isin(['plano','canal','faixa_tempo'])].copy();view['churn_ponderado']=view.churn_ponderado.map(pct);table(view)
text('Os recortes completos de coorte trimestral, canal de aquisição e evolução mensal estão na planilha. Os denominadores são exposições usuário-mês, não usuários únicos. Comparações entre coortes de idades diferentes têm exposição e calendário distintos; não são estimativas ajustadas.')
text('A faixa de 181 a 365 dias apresenta aumento de 6,46% para 9,16% em jan-set. Influenciador tem taxa de 9,00% em jan-set/2026 contra 7,57% no mesmo período de 2025. São pistas para investigar renovação anual e fim do desconto de primeiro ano, não provas desses mecanismos.')
text('A piora está presente em ambos os planos e em todos os canais observáveis. Não atribuir a subida apenas à mudança de composição ou ao reajuste sem análise ajustada. O anual mantém menor taxa por desenho do ciclo; analisar risco condicionado à renovação na etapa seguinte.')
h('## 7. Sinais anteriores à saída')
text('Para cada mês de churn, uso, recusas e tickets são medidos no mês calendário anterior. A queda de uso compara esse mês com o anterior a ele. Não usa uso do mês da saída nem dados posteriores à referência. Janeiro/2025 não tem uso de dezembro/2024 no anexo: fica sem registro, não zero.')
for key,g in signal_tables.items():
 text(key);v=g.copy();v['churn']=v.churn.map(pct);table(v)
text('As taxas são agregações descritivas sobre o painel repetido. Não foram usados testes ingênuos de independência entre todas as linhas, pois o mesmo usuário aparece em vários meses. Falta ajustar plano, tempo de casa, sazonalidade e campanha. Recusa anterior é sinal de risco, não classificação automática de churn involuntário. A aba de fins por motivo separa os eventos explicitamente marcados como falha_pagamento; os demais não são automaticamente voluntários.')
text('A forte associação de suporte recente pode refletir pedidos de cancelamento e problemas de cobrança, não necessariamente falha técnica. Renovação desligada é sinal já conhecido na referência e operacionalmente relevante, mas representa intenção declarada; convém avaliar modelos separados antes/depois desse aviso.')
h('## 8. Primeira avaliação da campanha de agosto')
text('Estimando: diferença de incidência de expiração entre grupos sorteados (intenção de tratar). Janela principal operacional deste relatório: após 10/08/2026 10h até 09/09/2026 10h, 30 dias. Não foi pré-especificada antes de ver os dados; conclusões são exploratórias. Aberturas e cliques não definem os grupos de análise.')
v=cg.copy();v['churn']=v.churn.map(pct);table(v)
text('Efeitos: valores positivos significam redução de churn em relação ao controle; IC95% por aproximação normal para duas proporções independentes.');table(effect[['grupo','reducao_churn_pp','ic95_inferior_pp','ic95_superior_pp','p_valor','p_bonferroni']].round(4))
text('Os três intervalos em 30 dias incluem zero. Os p-valores ajustados por Bonferroni para os três tratamentos também não indicam diferença. Ausência de significância não prova ausência de efeito.')
text('Análise secundária até o corte e sensibilidade restrita a elegíveis:');table(secondary[['analise','grupo','churn','reducao_pp','ic95_inf_pp','ic95_sup_pp','p_valor']].round(4))
text('No horizonte até o corte, cupom: p=0,0044; Bonferroni apenas entre os três braços desse horizonte ≈0,0133. Isso não corrige toda a exploração de horizontes e subgrupos. O resultado não deve ser apresentado como confirmação pré-registrada. No horizonte de 30 dias, excluir inconsistências de elegibilidade não muda a conclusão.')
text('Balanceamento prévio por grupo:');table(balance.round(4));table(balance_tests.round(4))
text('Plano e canal não mostram desequilíbrio grosseiro nos testes apresentados. Isso não prova balanceamento de todas as variáveis. Há 41 planos históricos desconhecidos no momento do envio, por falta de evento de início anterior ao horário.')
text('Inconsistência temporal: 48 dos 14.120 sorteados não estavam em ciclo pago vigente exatamente às 10h segundo os timestamps, e 1 pagante elegível não aparece na lista. O arquivo de campanha parece usar convenção de data diferente da hora real. Preservou-se a atribuição original no ITT; a sensibilidade exclui os 48. Validar a extração antes de uso operacional. A granularidade por dia usada na série mensal não deve ser confundida com elegibilidade instantânea.')
text('Por plano, os resultados de 30 dias estão na planilha e são exploratórios, sem teste confirmatório de heterogeneidade. A maior redução pontual de cupom no mensal não prova efeito diferencial. Não selecionar subgrupos apenas por significância.')
h('## 9. Retorno: o que é e o que não é possível concluir')
text('Não há ROI realizado estimado nesta etapa. Falta medir resgate e aplicação das ofertas, calendário de receita preservada, renovação anual, custo de IA e efeito ao longo do horizonte econômico. O efeito de push+e-mail não identifica o efeito de WhatsApp ou CS.')
text('A aba economia_cenarios é apenas uma análise de sensibilidade, não recomendação de gasto: retidos incrementais de 30 dias × margem mensal média aproximada − e-mails − oferta integral se todos os cupons enviados fossem concedidos. Mensal a R$39,90, anual a R$299,90/12, taxas do enunciado e servir a R$4,20. Ignora descontos anuais de aquisição, momento de cobrança e permanência futura. Esse cenário conservador cobra R$35,91 por mensal e R$25 por anual no braço cupom; não é custo observado nem fluxo mensal de caixa. A incerteza do efeito experimental precisa ser propagada antes de decidir.')
text('Os limites de R$40 mil/mês, 1.500 WhatsApps e 300 atendimentos de CS ainda não foram otimizados. A etapa seguinte deve priorizar valor incremental líquido, não apenas risco ou quantidade de retidos.')
h('## 10. Hipóteses, limitações e próximos passos')
text('Verificado: aumento de churn; base reconciliada; associações de uso/recusa/suporte; diferenças experimentais por horizonte. Hipóteses: sensibilidade ao reajuste; renovação de anuais com desconto de influenciador; efeitos técnicos; alterações de perfil de aquisição. Ainda não verificado: contribuição causal de cada hipótese, ganho de texto/imagem, desempenho preditivo, uplift individual e lucro de uma política.')
text('Prioridades: (1) validar convenções temporais da campanha e carência; (2) padronizar e investigar treinos inválidos; (3) ajustar diagnóstico por plano, renovação, coorte e sazonalidade; (4) criar cortes históricos com alvo de 60 dias completo e validação temporal; (5) treinar referência logística e comparar árvores/ensembles; (6) avaliar extração de texto com amostra rotulada e adquirir imagens; (7) validar política econômica em novo experimento.')
text('Não existe desfecho observado de 30/09 a 29/11/2026 nesta base. Não é possível medir desempenho real nessa janela. Dados atuais de perfil, curtidas acumuladas e status final de suporte podem vazar informação em cortes históricos e exigem tratamento específico.')
h('## 11. Como usar os produtos')
text('A planilha reúne auditoria, conciliação, série, recortes, sinais e campanha. Percentuais de colunas churn/pct são frações: 0,05 significa 5%; efeitos em pp já estão em pontos percentuais. Datas e nomes de campos estão preservados no código. Arquivos Parquet exigem pandas com pyarrow ou ferramenta equivalente.')
text('historico_mensal: uma linha por pagante presente no início do mês. historico_situacao_eventos: uma linha por evento mais saldo inicial legado em 01/07/2024. intervalos: início/fim por assinatura paga. sinais_anteriores: painel mensal enriquecido apenas com informação anterior. campanha_resultados_individuais: alocação original, atributos prévios, elegibilidade e desfecho de 30 dias.')
text('Reprodução: Python com pandas, numpy, scipy, matplotlib, pyarrow, openpyxl e python-docx. Colocar o script junto aos anexos originais e executá-lo. Os produtos não substituem os dados brutos e não modificam os anexos. O código usa atributos atuais somente para identificar planos pagos/cortesias; não usa plano atual para atribuir plano histórico.')
report='\n'.join(lines);(OUT/'relatorio_diagnostico.md').write_text(report,encoding='utf-8')
doc=Document();doc.core_properties.title='Vitaliza — Diagnóstico inicial';doc.core_properties.author='Julius'
from docx.shared import Pt
from docx.oxml import OxmlElement
doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
i=0
while i<len(lines):
 line=lines[i]
 if ' | ' in line:
  headers=line.split(' | ');t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
  for c,val in zip(t.rows[0].cells,headers):c.text=val.replace('_',' ')
  rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
  i+=1
  while i<len(lines) and ' | ' in lines[i]:
   for c,val in zip(t.add_row().cells,lines[i].split(' | ')):c.text=val
   i+=1
  for row in t.rows:
   for cell in row.cells:
    for para in cell.paragraphs:
     for run in para.runs:run.font.size=Pt(7)
  continue
 if line.startswith('# '):doc.add_heading(line[2:],0)
 elif line.startswith('## '):doc.add_heading(line[3:],1)
 elif line:doc.add_paragraph(line)
 i+=1
for image in ['01_churn_mensal.png','02_churn_plano.png','03_campanha.png']:
 doc.add_page_break();doc.add_picture(str(OUT/image),width=Inches(6.3))
doc.save(OUT/'relatorio_diagnostico.docx')
import shutil
shutil.copyfile(__file__,OUT/'produzir_entregas_vitaliza.py')
(OUT/'requirements.txt').write_text('pandas\nnumpy\nscipy\nmatplotlib\npyarrow\nopenpyxl\npython-docx\n')
with zipfile.ZipFile('entregas_vitaliza_completas.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in OUT.iterdir():z.write(p,arcname=str(p))
print('ENTREGAS FINALIZADAS')
